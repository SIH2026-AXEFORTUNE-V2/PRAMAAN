"""Source ingestion: normalises text, documents, images, audio/video and URLs.

Textual formats are converted to plain text locally. PDFs are text-extracted; PDFs
without a text layer, images, audio and video are kept as media for the perception
stage (vision model / speech-to-text).
"""
from __future__ import annotations

import hashlib
import io
import mimetypes
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone

import httpx
from bs4 import BeautifulSoup

from . import config

TEXT_EXT = {".txt", ".md", ".markdown", ".csv", ".json", ".log", ".xml", ".yaml", ".yml", ".rtf"}
HTML_EXT = {".html", ".htm"}
MEDIA_PREFIX = ("image/", "audio/", "video/")
YOUTUBE_RE = re.compile(r"^https?://(www\.)?(youtube\.com/watch\?v=|youtu\.be/|youtube\.com/shorts/)", re.I)


class IngestError(ValueError):
    pass


@dataclass
class Source:
    kind: str                       # text | document | image | audio | video | web
    name: str
    mime: str = "text/plain"
    text: str = ""                  # extracted text (may be empty for pure media)
    media: bytes | None = None      # raw bytes for the perception stage
    url: str = ""
    notes: list[str] = field(default_factory=list)
    id: str = "S1"
    size_bytes: int = 0
    sha256: str = ""                # hash of the original bytes (provenance)
    pages: list[list[int]] = field(default_factory=list)   # [start, end) char offsets of each PDF page in text
    pdf: bytes | None = None        # original PDF bytes, kept in memory for evidence bounding boxes
    uploaded_at: str = field(default_factory=lambda: datetime.now(timezone.utc).isoformat(timespec="seconds"))

    def describe(self) -> dict:
        return {
            "id": self.id, "kind": self.kind, "name": self.name, "mime": self.mime, "url": self.url,
            "size_bytes": self.size_bytes, "sha256": self.sha256, "uploaded_at": self.uploaded_at,
            "chars": len(self.text), "words": len(self.text.split()), "page_count": len(self.pages) or None,
            "has_media": self.media is not None, "notes": self.notes, "status": "received",
        }

    def page_of(self, offset: int) -> int | None:
        for i, (a, b) in enumerate(self.pages, 1):
            if a <= offset < b:
                return i
        return None


def _hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _clip(text: str) -> str:
    text = re.sub(r"\n{3,}", "\n\n", text.replace("\r\n", "\n")).strip()
    return text[: config.MAX_SOURCE_CHARS]


def _html_to_text(html: str) -> tuple[str, str]:
    soup = BeautifulSoup(html, "html.parser")
    title = (soup.title.string or "").strip() if soup.title and soup.title.string else ""
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "aside", "form", "svg", "iframe"]):
        tag.decompose()
    root = soup.find("article") or soup.find("main") or soup.body or soup
    blocks = [el.get_text(" ", strip=True) for el in root.find_all(["h1", "h2", "h3", "h4", "p", "li", "blockquote", "td"])]
    text = "\n".join(b for b in blocks if len(b) > 1)
    if len(text) < 200:
        text = root.get_text("\n", strip=True)
    return title, text


def from_text(text: str) -> Source:
    text = _clip(text or "")
    if len(text) < 3:
        raise IngestError("Please provide some source content.")
    first = text.split("\n", 1)[0][:80]
    raw = text.encode()
    return Source(kind="text", name=first or "Pasted text", text=text, size_bytes=len(raw), sha256=_hash(raw))


def from_file(filename: str, data: bytes, content_type: str | None) -> Source:
    src = _from_file(filename, data, content_type)
    src.size_bytes, src.sha256 = len(data), _hash(data)
    return src


def _pdf_text(data: bytes) -> tuple[str, list[list[int]]]:
    """Extract text page by page and remember where each page starts, so claims map to real page numbers."""
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    text, pages = "", []
    for page in reader.pages:
        chunk = re.sub(r"[ \t]+\n", "\n", page.extract_text() or "").strip()
        start = len(text)
        text += chunk + "\n\n"
        pages.append([start, len(text)])
    return text, pages


def _from_file(filename: str, data: bytes, content_type: str | None) -> Source:
    if not data:
        raise IngestError("The uploaded file is empty.")
    if len(data) > config.MAX_UPLOAD_MB * 1024 * 1024:
        raise IngestError(f"File exceeds the {config.MAX_UPLOAD_MB} MB limit.")
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    mime = (content_type if content_type and content_type != "application/octet-stream" else None) \
        or mimetypes.guess_type(filename)[0] or "application/octet-stream"

    if ext in TEXT_EXT or mime.startswith("text/plain"):
        return Source(kind="document", name=filename, mime="text/plain", text=_clip(data.decode("utf-8", "replace")))
    if ext in HTML_EXT or mime == "text/html":
        title, text = _html_to_text(data.decode("utf-8", "replace"))
        return Source(kind="document", name=title or filename, mime="text/html", text=_clip(text))
    if ext == ".pdf" or mime == "application/pdf":
        text, pages = "", []
        try:
            text, pages = _pdf_text(data)
        except Exception:  # noqa: BLE001 - scanned or unusual PDFs fall back to vision
            pass
        if len(text.strip()) < 50:
            src = Source(kind="document", name=filename, mime="application/pdf", media=data)
            src.notes.append("No text layer found; pages will be read by the vision model.")
            return src
        clipped = text[: config.MAX_SOURCE_CHARS]
        pages = [[a, min(b, len(clipped))] for a, b in pages if a < len(clipped)]
        return Source(kind="document", name=filename, mime="application/pdf", text=clipped, pages=pages, pdf=data)
    if ext == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return Source(kind="document", name=filename, mime=mime, text=_clip("\n".join(parts)))
    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            parts.append(f"[Slide {i}]")
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    parts.append(shape.text_frame.text)
        return Source(kind="document", name=filename, mime=mime, text=_clip("\n".join(parts)))
    if mime.startswith(MEDIA_PREFIX):
        kind = mime.split("/")[0]
        return Source(kind=kind, name=filename, mime=mime, media=data)
    raise IngestError(f"Unsupported file type: {filename} ({mime}). "
                      "Use text, PDF, DOCX, PPTX, HTML, images, audio or video.")


async def from_url(url: str) -> Source:
    url = url.strip()
    if not re.match(r"^https?://", url, re.I):
        raise IngestError("URL must start with http:// or https://")
    if YOUTUBE_RE.match(url):
        raise IngestError("YouTube links cannot be downloaded here. Upload the video file or paste its transcript.")
    try:
        async with httpx.AsyncClient(follow_redirects=True, timeout=20,
                                     headers={"User-Agent": "Mozilla/5.0 (PRAMAAN source ingestor)"}) as client:
            resp = await client.get(url)
            resp.raise_for_status()
    except httpx.HTTPError as e:
        raise IngestError(f"Could not fetch the URL: {e}") from e
    ctype = resp.headers.get("content-type", "").split(";")[0].strip()
    if ctype == "application/pdf" or ctype.startswith(MEDIA_PREFIX):
        src = from_file(url.rsplit("/", 1)[-1] or "download", resp.content, ctype)
        src.url = url
        return src
    title, text = _html_to_text(resp.text)
    if len(text) < 100:
        raise IngestError("The page did not contain readable article text.")
    raw = resp.content
    return Source(kind="web", name=title or url, mime="text/html", text=_clip(text), url=url,
                  size_bytes=len(raw), sha256=_hash(raw))
