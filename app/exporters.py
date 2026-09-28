"""Export deliverables to Markdown, plain text, SRT, DOCX, PPTX, SVG and a ZIP bundle."""
from __future__ import annotations

import html
import io
import json
import re
import textwrap
import zipfile

from .catalog import OUTPUT_TYPES

# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------


def _bullets(items, fmt=lambda x: x) -> str:
    return "\n".join(f"- {fmt(i)}" for i in items or [])


def to_markdown(otype: str, d: dict) -> str:
    if otype == "video":
        scenes = "\n\n".join(
            f"### Scene {s['scene_number']} · {s['start_sec']:.0f}s-{s['end_sec']:.0f}s · {s['title']}\n"
            f"- **Visual:** {s['visual_description']}\n- **Camera:** {s['camera_direction']}\n"
            f"- **On-screen text:** {s['on_screen_text'] or '—'}\n- **Narration:** {s['narration']}\n"
            f"- **Visual recommendation:** {s['visual_recommendation']}\n- **Audio:** {s['audio_cue']}"
            for s in d.get("scenes", []))
        vs = d.get("visual_style", {})
        subs = "\n".join(f"{s['index']}. [{s['start']} → {s['end']}] {s['text']}" for s in d.get("subtitles", []))
        return (f"# {d['title']}\n\n_{d['logline']}_\n\n**Platform:** {d['target_platform']} · **Length:** {d['duration_seconds']}s\n\n"
                f"## Script\n\n{d['script']}\n\n## Storyboard\n\n{scenes}\n\n## Visual direction\n\n"
                f"- **Look & feel:** {vs.get('look_and_feel', '')}\n- **Palette:** {', '.join(vs.get('color_palette', []))}\n"
                f"- **Typography:** {vs.get('typography', '')}\n- **Music:** {vs.get('music_direction', '')}\n\n"
                f"## Subtitles\n\n{subs}\n\n**Call to action:** {d['call_to_action']}\n")
    if otype == "linkedin":
        return (f"{d['post']}\n\n{' '.join(d.get('hashtags', []))}\n\n---\n**Alternate hooks**\n{_bullets(d.get('alternate_hooks'))}\n\n"
                f"**Suggested visual:** {d['suggested_visual']}\n\n**Best time to post:** {d['best_time_to_post']}\n")
    if otype == "twitter":
        tw = "\n\n".join(t["text"] for t in d.get("tweets", []))
        alt = _bullets(d.get("alternates"))
        return f"{tw}\n\n---\n**Alternates**\n{alt}\n\n**Suggested media:** {d['suggested_media']}\n"
    if otype == "advisory":
        recs = "\n".join(f"| {r['priority']} | {r['action']} | {r['owner']} |" for r in d.get("recommendations", []))
        det = "\n\n".join(f"### {x['heading']}\n{x['body']}" for x in d.get("details", []))
        ind = "\n".join(f"| {i['type']} | `{i['value']}` |" for i in d.get("indicators", []))
        out = (f"# {d['title']}\n\n**{d['advisory_id']}** · {d['classification']} · Severity: **{d['severity'].upper()}** · Status: {d['status']}\n\n"
               f"## Summary\n{d['summary']}\n\n## Background\n{d['background']}\n\n## Affected\n{_bullets(d.get('affected'))}\n\n"
               f"## Impact\n{d['impact']}\n\n## Details\n{det}\n\n## Recommendations\n| Priority | Action | Owner |\n|---|---|---|\n{recs}\n")
        if ind:
            out += f"\n## Indicators\n| Type | Value |\n|---|---|\n{ind}\n"
        if d.get("references"):
            out += f"\n## References\n{_bullets(d['references'])}\n"
        return out + f"\n**Contact:** {d['contact']}\n"
    if otype == "infographic":
        secs = "\n".join(f"- **{s['heading']}** ({s['icon']}) {s['stat_value']} {s['stat_label']} — {s['text']}" for s in d.get("sections", []))
        lay = d.get("layout", {})
        chart = d.get("chart", {})
        cdata = ", ".join(f"{x['label']}: {x['value']}{chart.get('unit', '')}" for x in chart.get("data", []))
        return (f"# {d['title']}\n_{d['subtitle']}_\n\n**Headline stat:** {d['headline_stat']['value']} — {d['headline_stat']['label']}\n\n"
                f"## Sections\n{secs}\n\n## Chart\n{chart.get('title', '')} ({chart.get('chart_type')}): {cdata or '—'}\n\n"
                f"## Key messages\n{_bullets(d.get('key_messages'))}\n\n## Layout recommendation\n- **Type:** {lay.get('type')}\n"
                f"- **Why:** {lay.get('rationale')}\n- **Palette:** {', '.join(lay.get('color_palette', []))}\n- **Typography:** {lay.get('typography')}\n"
                f"- **Reading path:** {' → '.join(lay.get('visual_hierarchy', []))}\n\n**CTA:** {d['call_to_action']}\n\n_{d['source_note']}_\n")
    if otype == "executive_summary":
        acts = "\n".join(f"| {a['action']} | {a['owner']} | {a['timeline']} |" for a in d.get("recommended_actions", []))
        mets = " · ".join(f"**{m['label']}:** {m['value']}" for m in d.get("metrics", []))
        return (f"# {d['title']}\n\n> **Bottom line:** {d['bottom_line']}\n\n## Situation\n{d['situation']}\n\n"
                + (f"{mets}\n\n" if mets else "")
                + f"## Key findings\n{_bullets(d.get('key_findings'))}\n\n## Implications\n{_bullets(d.get('implications'))}\n\n"
                f"## Recommended actions\n| Action | Owner | Timeline |\n|---|---|---|\n{acts}\n\n"
                + (f"## Risks\n{_bullets(d['risks'])}\n\n" if d.get("risks") else "")
                + f"**Decision required:** {d['decision_required']}\n")
    if otype == "presentation":
        parts = [f"# {d['title']}\n_{d['subtitle']}_\n"]
        for i, s in enumerate(d.get("slides", []), 1):
            body = _bullets(s.get("bullets"))
            if s.get("right_bullets"):
                body += "\n\n" + _bullets(s["right_bullets"])
            if s.get("stat_value"):
                body += f"\n\n**{s['stat_value']}** — {s.get('stat_label', '')}"
            if s.get("quote"):
                body += f"\n\n> {s['quote']}"
            parts.append(f"## Slide {i} · {s['title']}  `{s['layout']}`\n{body}\n\n*Visual:* {s.get('visual_suggestion', '')}\n\n"
                         f"**Speaker notes:** {s.get('speaker_notes', '')}\n")
        return "\n".join(parts)
    return "```json\n" + json.dumps(d, indent=2, ensure_ascii=False) + "\n```"


def to_text(otype: str, d: dict) -> str:
    if otype == "linkedin":
        return d["post"] + "\n\n" + " ".join(d.get("hashtags", []))
    if otype == "twitter":
        return "\n\n".join(t["text"] for t in d.get("tweets", []))
    return re.sub(r"[*_`#>|]", "", to_markdown(otype, d))


def to_srt(d: dict) -> str:
    return "\n".join(f"{s['index']}\n{s['start']} --> {s['end']}\n{s['text']}\n" for s in d.get("subtitles", []))


# ---------------------------------------------------------------------------
# DOCX (advisory, executive summary, or any markdown)
# ---------------------------------------------------------------------------


def to_docx(otype: str, d: dict, footer_text: str = "") -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    lines = to_markdown(otype, d).splitlines()
    table_rows: list[list[str]] = []

    def flush_table():
        nonlocal table_rows
        rows = [r for r in table_rows if not all(set(c) <= set("-: ") for c in r)]
        if rows:
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Light Grid Accent 1"
            for i, r in enumerate(rows):
                for j, c in enumerate(r[: len(rows[0])]):
                    t.cell(i, j).text = c.strip().strip("`")
        table_rows = []

    def add_runs(par, text):
        for chunk in re.split(r"(\*\*[^*]+\*\*)", text):
            if chunk.startswith("**") and chunk.endswith("**"):
                par.add_run(chunk[2:-2]).bold = True
            elif chunk:
                par.add_run(chunk.replace("_", "") if chunk.startswith("_") else chunk)

    for line in lines:
        if line.startswith("|"):
            table_rows.append([c for c in line.strip("|").split("|")])
            continue
        flush_table()
        if line.startswith("# "):
            h = doc.add_heading(line[2:], level=0)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x0B, 0x0B, 0x0C)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            add_runs(doc.add_paragraph(style="List Bullet"), line[2:])
        elif line.startswith("> "):
            p = doc.add_paragraph(style="Intense Quote")
            add_runs(p, line[2:])
        elif line.strip() and line.strip() != "---":
            add_runs(doc.add_paragraph(), line)
    flush_table()
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = footer_text or "PRAMAAN · AI Transformation Workspace"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PPTX (presentation)
# ---------------------------------------------------------------------------


def to_pptx(d: dict, images: dict[int, bytes] | None = None) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu, Inches, Pt

    BLACK, WHITE, INK, MUTED, ACCENT, LINE = (RGBColor(0x0B, 0x0B, 0x0C), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x11, 0x18, 0x27),
                                              RGBColor(0x6B, 0x72, 0x80), RGBColor(0xF5, 0xA5, 0x24), RGBColor(0xE5, 0xE7, 0xEB))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    H = prs.slide_height

    def bg(slide, color):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def text(slide, x, y, w, h, value, size, color, bold=False, font="Inter", italic=False):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.text = value or ""
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(size), bold, italic, font
                r.font.color.rgb = color
        return tf

    def bullet_box(slide, x, y, w, h, items, color):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = f"■  {item}"
            r.font.size, r.font.name = Pt(20), "Inter"
            r.font.color.rgb = color
            p.space_after = Pt(14)

    def label(slide, value, dark):
        text(slide, Inches(0.7), Inches(0.45), Inches(8), Inches(0.4), value.upper(), 11, MUTED if not dark else RGBColor(0x9C, 0xA3, 0xAF), bold=True)

    slides = d.get("slides", [])
    for idx, s in enumerate(slides, 1):
        slide = prs.slides.add_slide(blank)
        lay = s.get("layout", "bullets")
        dark = lay in ("title", "closing", "big_stat")
        bg(slide, BLACK if dark else WHITE)
        fg = WHITE if dark else INK
        label(slide, f"{d.get('title', '')[:60]}  ·  {idx:02d}", dark)
        img = (images or {}).get(idx)
        if lay == "title":
            tw = Inches(6.0) if img else Inches(11.5)
            if img:
                slide.shapes.add_picture(io.BytesIO(img), Inches(7.333), 0, Inches(6.0), H)
                text(slide, Inches(7.5), H - Inches(0.55), Inches(5.6), Inches(0.3), "AI-generated illustration", 9,
                     RGBColor(0xD1, 0xD5, 0xDB))
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(2.4), Inches(1.2), Emu(60000))
            bar.fill.solid()
            bar.fill.fore_color.rgb = ACCENT
            bar.line.fill.background()
            text(slide, Inches(0.7), Inches(2.6), tw, Inches(2.2), s["title"], 44 if img else 48, WHITE, bold=True, font="Georgia")
            text(slide, Inches(0.7), Inches(4.9), tw, Inches(1), d.get("subtitle", ""), 22, RGBColor(0xD1, 0xD5, 0xDB))
        elif lay == "big_stat":
            text(slide, Inches(0.7), Inches(1.3), Inches(12), Inches(1), s["title"], 30, WHITE, bold=True, font="Georgia")
            text(slide, Inches(0.7), Inches(2.5), Inches(12), Inches(2.4), s.get("stat_value") or "", 110, ACCENT, bold=True, italic=True)
            text(slide, Inches(0.7), Inches(5.0), Inches(11), Inches(1), s.get("stat_label") or "", 24, RGBColor(0xD1, 0xD5, 0xDB))
        elif lay == "quote":
            text(slide, Inches(0.7), Inches(1.2), Inches(12), Inches(0.9), s["title"], 30, INK, bold=True, font="Georgia")
            text(slide, Inches(1.2), Inches(2.5), Inches(11), Inches(3), f"“{s.get('quote') or ''}”", 32, INK, italic=True, font="Georgia")
        else:
            text(slide, Inches(0.7), Inches(1.0), Inches(12), Inches(1.2), s["title"], 34, fg, bold=True, font="Georgia")
            if lay == "two_column":
                bullet_box(slide, Inches(0.7), Inches(2.5), Inches(5.8), Inches(4.2), s.get("bullets", []), fg)
                bullet_box(slide, Inches(6.9), Inches(2.5), Inches(5.8), Inches(4.2), s.get("right_bullets", []), fg)
            else:
                bullet_box(slide, Inches(0.7), Inches(2.5), Inches(8.2), Inches(4.3), s.get("bullets", []), fg)
                if img:
                    slide.shapes.add_picture(io.BytesIO(img), Inches(9.3), Inches(2.5), Inches(3.4), Inches(3.4 * 9 / 16))
                    text(slide, Inches(9.3), Inches(2.5 + 3.4 * 9 / 16 + 0.05), Inches(3.4), Inches(0.3), "AI-generated illustration", 9, MUTED)
                elif s.get("visual_suggestion") and lay != "closing":
                    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.3), Inches(2.5), Inches(3.4), Inches(3.6))
                    card.fill.solid()
                    card.fill.fore_color.rgb = RGBColor(0xF6, 0xF7, 0xF9) if not dark else RGBColor(0x1F, 0x20, 0x23)
                    card.line.color.rgb = LINE
                    card.adjustments[0] = 0.08
                    tf = card.text_frame
                    tf.word_wrap = True
                    tf.text = "VISUAL\n" + s["visual_suggestion"]
                    for p in tf.paragraphs:
                        for r in p.runs:
                            r.font.size, r.font.name = Pt(13), "Inter"
                            r.font.color.rgb = MUTED
        text(slide, Inches(0.7), H - Inches(0.6), Inches(6), Inches(0.3), "PRAMAAN", 10, MUTED, bold=True)
        if s.get("speaker_notes"):
            slide.notes_slide.notes_text_frame.text = s["speaker_notes"]
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# SVG infographic
# ---------------------------------------------------------------------------
ICON_PATHS = {
    "alert": "M12 3 2 21h20L12 3zm0 6v5m0 3v.5",
    "shield": "M12 3 4 6v6c0 5 3.5 8 8 9 4.5-1 8-4 8-9V6l-8-3z",
    "chart": "M4 20V10m6 10V4m6 16v-7m4 7H2",
    "users": "M16 20v-2a4 4 0 0 0-8 0v2M12 11a3 3 0 1 0 0-6 3 3 0 0 0 0 6zm8 9v-2a3 3 0 0 0-2-2.8M4 20v-2a3 3 0 0 1 2-2.8",
    "globe": "M12 21a9 9 0 1 0 0-18 9 9 0 0 0 0 18zM3 12h18M12 3c3 3 3 15 0 18M12 3c-3 3-3 15 0 18",
    "clock": "M12 21a9 9 0 1 0 0-18 9 9 0 0 0 0 18zm0-13v5l3 2",
    "money": "M3 7h18v10H3zM12 15a3 3 0 1 0 0-6 3 3 0 0 0 0 6z",
    "target": "M12 21a9 9 0 1 0 0-18 9 9 0 0 0 0 18zm0-4a5 5 0 1 0 0-10 5 5 0 0 0 0 10zm0-4a1 1 0 1 0 0-2 1 1 0 0 0 0 2z",
    "lightbulb": "M9 18h6M10 21h4M12 3a6 6 0 0 0-3.5 10.9V16h7v-2.1A6 6 0 0 0 12 3z",
    "lock": "M5 11h14v10H5zM8 11V7a4 4 0 0 1 8 0v4",
    "trend": "M3 17l6-6 4 4 8-8m0 0h-5m5 0v5",
    "check": "M12 21a9 9 0 1 0 0-18 9 9 0 0 0 0 18zm-4-9 3 3 5-6",
    "doc": "M6 3h9l4 4v14H6zM14 3v5h5M9 13h6M9 17h6",
    "server": "M4 4h16v6H4zM4 14h16v6H4zM8 7h.01M8 17h.01",
}


def _wrap(text: str, width: int) -> list[str]:
    return textwrap.wrap(text or "", width=width) or [""]


def to_svg(d: dict) -> str:
    pal = (d.get("layout", {}).get("color_palette") or []) + ["#0B0B0C", "#E5484D", "#F5A524", "#16A34A", "#2563EB"]
    hexes = [c for c in pal if re.fullmatch(r"#[0-9A-Fa-f]{6}", c or "")]
    dark = min(hexes, key=lambda c: sum(int(c[i:i + 2], 16) for i in (1, 3, 5)))
    accents = [c for c in hexes if c != dark and sum(int(c[i:i + 2], 16) for i in (1, 3, 5)) < 690][:3] or ["#F5A524"]
    e = html.escape
    W = 1080
    out: list[str] = []
    y = 0
    # header
    title_lines = _wrap(d.get("title", ""), 30)
    sub_lines = _wrap(d.get("subtitle", ""), 62)
    head_h = 120 + 64 * len(title_lines) + 34 * len(sub_lines) + 40
    out.append(f'<rect x="0" y="0" width="{W}" height="{head_h}" fill="{dark}"/>')
    out.append(f'<text x="64" y="84" class="lbl" fill="{accents[0]}">INFOGRAPHIC</text>')
    ty = 150
    for ln in title_lines:
        out.append(f'<text x="64" y="{ty}" class="h1" fill="#fff">{e(ln)}</text>')
        ty += 64
    for ln in sub_lines:
        out.append(f'<text x="64" y="{ty}" class="sub" fill="#D1D5DB">{e(ln)}</text>')
        ty += 34
    y = head_h + 56
    # headline stat
    hs = d.get("headline_stat", {})
    out.append(f'<rect x="64" y="{y}" width="{W - 128}" height="170" rx="20" fill="#F6F7F9" stroke="#E5E7EB"/>')
    out.append(f'<text x="104" y="{y + 115}" class="stat" fill="{accents[0]}">{e(hs.get("value", ""))}</text>')
    for i, ln in enumerate(_wrap(hs.get("label", ""), 30)[:3]):
        out.append(f'<text x="520" y="{y + 72 + i * 34}" class="statlbl" fill="#111827">{e(ln)}</text>')
    y += 170 + 48
    # sections grid (2 columns)
    secs = d.get("sections", [])
    col_w = (W - 128 - 32) // 2
    for i in range(0, len(secs), 2):
        row = secs[i:i + 2]
        heights = []
        for s in row:
            heights.append(150 + 30 * len(_wrap(s.get("text", ""), 34)))
        rh = max(heights)
        for j, s in enumerate(row):
            x = 64 + j * (col_w + 32)
            color = accents[(i + j) % len(accents)]
            out.append(f'<rect x="{x}" y="{y}" width="{col_w}" height="{rh}" rx="18" fill="#fff" stroke="#E5E7EB"/>')
            out.append(f'<circle cx="{x + 52}" cy="{y + 54}" r="28" fill="{color}" fill-opacity="0.12"/>')
            out.append(f'<g transform="translate({x + 38},{y + 40}) scale(1.17)"><path d="{ICON_PATHS.get(s.get("icon"), ICON_PATHS["doc"])}" '
                       f'fill="none" stroke="{color}" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"/></g>')
            out.append(f'<text x="{x + 96}" y="{y + 50}" class="lbl" fill="#6B7280">{e(s.get("heading", "").upper())}</text>')
            if s.get("stat_value"):
                out.append(f'<text x="{x + 96}" y="{y + 88}" class="mstat" fill="#0B0B0C">{e(s["stat_value"])}'
                           f'<tspan class="small" fill="#6B7280" dx="10">{e(s.get("stat_label", ""))}</tspan></text>')
            ty = y + 130
            for ln in _wrap(s.get("text", ""), 34):
                out.append(f'<text x="{x + 32}" y="{ty}" class="body" fill="#374151">{e(ln)}</text>')
                ty += 30
        y += rh + 32
    # chart
    chart = d.get("chart") or {}
    data = [p for p in chart.get("data", []) if isinstance(p.get("value"), (int, float))]
    if chart.get("chart_type") == "bar" and len(data) >= 2:
        y += 8
        out.append(f'<text x="64" y="{y + 20}" class="lbl" fill="#6B7280">{e((chart.get("title") or "").upper())}</text>')
        y += 48
        mx = max(abs(p["value"]) for p in data) or 1
        for k, p in enumerate(data):
            bw = int((W - 128 - 260 - 140) * abs(p["value"]) / mx)
            out.append(f'<text x="64" y="{y + 28}" class="body" fill="#111827">{e(str(p.get("label", ""))[:22])}</text>')
            out.append(f'<rect x="324" y="{y + 6}" width="{max(bw, 4)}" height="30" rx="6" fill="{dark if k == 0 else accents[k % len(accents)]}"/>')
            val = f'{p["value"]:g}{chart.get("unit", "")}'
            out.append(f'<text x="{324 + max(bw, 4) + 14}" y="{y + 29}" class="mono" fill="#111827">{e(val)}</text>')
            y += 48
        y += 24
    # key messages
    msgs = d.get("key_messages", [])
    if msgs:
        box_h = 70 + sum(34 * len(_wrap(m, 58)) + 14 for m in msgs)
        out.append(f'<rect x="64" y="{y}" width="{W - 128}" height="{box_h}" rx="18" fill="{dark}"/>')
        out.append(f'<text x="104" y="{y + 50}" class="lbl" fill="{accents[0]}">KEY MESSAGES</text>')
        ty = y + 94
        for m in msgs:
            out.append(f'<circle cx="112" cy="{ty - 8}" r="5" fill="{accents[0]}"/>')
            for ln in _wrap(m, 58):
                out.append(f'<text x="134" y="{ty}" class="msg" fill="#fff">{e(ln)}</text>')
                ty += 34
            ty += 14
        y += box_h + 40
    out.append(f'<text x="64" y="{y + 10}" class="cta" fill="#0B0B0C">{e(d.get("call_to_action", ""))}</text>')
    out.append(f'<text x="64" y="{y + 46}" class="small" fill="#6B7280">{e(d.get("source_note", ""))}</text>')
    H = y + 90
    style = ("<style>text{font-family:Inter,'Helvetica Neue',Arial,sans-serif}.h1{font-family:Georgia,'Source Serif 4',serif;font-size:54px;font-weight:700}"
             ".sub{font-size:24px}.lbl{font-size:15px;font-weight:700;letter-spacing:2.5px}.stat{font-size:92px;font-weight:800;font-style:italic}"
             ".statlbl{font-size:28px;font-weight:600}.mstat{font-size:34px;font-weight:800}.small{font-size:16px;font-weight:500}"
             ".body{font-size:21px}.msg{font-size:24px;font-weight:500}.cta{font-size:28px;font-weight:800}"
             ".mono{font-family:'JetBrains Mono',Menlo,monospace;font-size:20px;font-weight:700}</style>")
    return (f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}">{style}'
            f'<rect width="{W}" height="{H}" fill="#FFFFFF"/>{"".join(out)}</svg>')


# ---------------------------------------------------------------------------
# Dispatch + bundle
# ---------------------------------------------------------------------------
MIME = {
    "mp4": "video/mp4",
    "md": "text/markdown; charset=utf-8", "txt": "text/plain; charset=utf-8", "json": "application/json",
    "srt": "application/x-subrip", "svg": "image/svg+xml",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


def render(otype: str, fmt: str, d: dict, provenance: str = "", images: dict[int, bytes] | None = None) -> bytes:
    """provenance: one-line lineage footer (transformation, version, output hash, approval) added to text formats."""
    tail = f"\n\n---\n{provenance}\n" if provenance else ""
    if fmt == "md":
        return (to_markdown(otype, d) + tail).encode()
    if fmt == "txt":
        return (to_text(otype, d) + tail).encode()
    if fmt == "json":
        return json.dumps({"content": d, "provenance": provenance} if provenance else d, indent=2, ensure_ascii=False).encode()
    if fmt == "srt" and otype == "video":
        return to_srt(d).encode()
    if fmt == "svg" and otype == "infographic":
        return to_svg(d).encode()
    if fmt == "docx":
        return to_docx(otype, d, provenance)
    if fmt == "pptx" and otype == "presentation":
        return to_pptx(d, images)
    raise ValueError(f"Format '{fmt}' is not available for {otype}.")


def bundle(t: dict, footer, image_files=None, attachments=None) -> bytes:
    """Approved artefacts (released versions) + evidence ledger + provenance manifest."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        manifest = {"transformation_id": t["id"], "title": t["title"], "contract": t["contract"],
                    "provenance": t["provenance"], "artifacts": {}}
        for otype, a in t["artifacts"].items():
            if a["approval"]["status"] != "approved":
                continue
            manifest["artifacts"][otype] = {"version": a["version"], "output_hash": a["output_hash"],
                                            "approved_by": a["approval"]["by"], "approved_at": a["approval"]["at"]}
            files = image_files(otype) if image_files else {}
            slide_imgs = {int(k.split("-")[1]): v for k, v in files.items() if k.startswith("slide-")}
            for fmt in OUTPUT_TYPES[otype]["exports"]:
                try:
                    z.writestr(f"{otype}/{otype}.{fmt}", render(otype, fmt, a["released"], footer(a), slide_imgs))
                except ValueError:
                    pass
            for slot, data in files.items():
                z.writestr(f"{otype}/illustrations/{slot}.jpg", data)
            for name, data in (attachments(otype) if attachments else {}).items():
                z.writestr(name, data)
        z.writestr("evidence_ledger.json", json.dumps(
            [{k: c[k] for k in ("claim_id", "label", "display_value", "modality", "source_name", "page", "status", "version")}
             for c in t["claims"]], indent=2, ensure_ascii=False))
        z.writestr("manifest.json", json.dumps(manifest, indent=2, ensure_ascii=False))
    return buf.getvalue()


def filename(t: dict, otype: str, fmt: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", (t.get("title") or "pramaan").lower()).strip("-")[:40] or "pramaan"
    return f"{t['id']}-{slug}-{otype}.{fmt}"
